import os
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app import models
from app.config import settings
from app.schemas import LedgerResponse

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(size=10, italic=True, color="666666")
BOLD = Font(bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="CCCCCC"),) * 4)
MONEY_FMT = '#,##0.00 "TRY"'
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right")


def _money(value: float) -> str:
    return f"{value:,.2f} TRY"


def _write_section(ws, start_row: int, title: str, rows, debt_label: str, credit_label: str, balance_value: float) -> int:
    ws.cell(row=start_row, column=1, value=title).font = Font(size=12, bold=True, color="1F4E78")
    start_row += 1

    headers = ["Tarih", "Tür", "Referans", "Açıklama", "Borç", "Yapılan Ödeme", "Bakiye"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    start_row += 1

    row_type_labels = {"invoice": "Fatura", "payment": "Ödeme"}
    for row in rows:
        ws.cell(row=start_row, column=1, value=row.date.strftime("%d.%m.%Y"))
        ws.cell(row=start_row, column=2, value=row_type_labels.get(row.row_type, row.row_type))
        ws.cell(row=start_row, column=3, value=row.reference)
        ws.cell(row=start_row, column=4, value=row.description or "")
        ws.cell(row=start_row, column=5, value=row.debit or None)
        ws.cell(row=start_row, column=6, value=row.credit or None)
        ws.cell(row=start_row, column=7, value=row.running_balance)
        for col in range(1, 8):
            cell = ws.cell(row=start_row, column=col)
            cell.border = THIN_BORDER
            if col in (5, 6, 7):
                cell.number_format = MONEY_FMT
        start_row += 1

    # Always shown as a non-negative "what's still owed" figure; if payments
    # exceeded invoices the excess becomes its own clearly-labeled credit line
    # instead of a bare negative number under a "debt" heading.
    start_row += 1
    ws.cell(row=start_row, column=4, value=debt_label).font = BOLD
    debt_cell = ws.cell(row=start_row, column=7, value=_money(max(balance_value, 0.0)))
    debt_cell.font = BOLD
    debt_cell.alignment = RIGHT
    start_row += 1

    if balance_value < -0.005:
        ws.cell(row=start_row, column=4, value=credit_label).font = BOLD
        credit_cell = ws.cell(row=start_row, column=7, value=_money(abs(balance_value)))
        credit_cell.font = BOLD
        credit_cell.alignment = RIGHT
        start_row += 1

    return start_row + 2


def export_company_ledger(ledger: LedgerResponse) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cari Hesap Ekstresi"

    ws.cell(row=1, column=1, value="BEMAS TREYLER - Cari Hesap Ekstresi").font = TITLE_FONT
    ws.cell(row=2, column=1, value=f"Firma: {ledger.company.name}  |  Vergi No: {ledger.company.tax_id}").font = SUBTITLE_FONT
    ws.cell(row=3, column=1, value=f"Rapor Tarihi: {dt.date.today().strftime('%d.%m.%Y')}").font = SUBTITLE_FONT

    has_receivable = len(ledger.receivable_rows) > 0
    has_payable = len(ledger.payable_rows) > 0

    row = 5
    if has_receivable:
        row = _write_section(
            ws, row, "MÜŞTERİ HESABI (Satış Faturaları / Tahsilatlar)",
            ledger.receivable_rows,
            "Net Müşteri Borcu:",
            "Müşterinin Alacağı:",
            ledger.balance.net_customer_debt,
        )
    if has_payable:
        row = _write_section(
            ws, row, "TEDARİKÇİ HESABI (Alış Faturaları / Ödemeler)",
            ledger.payable_rows,
            "Net Firma Borcu:",
            "Bizim Alacağımız:",
            ledger.balance.net_company_debt,
        )

    # The overall figure only adds information when both sides are active -
    # otherwise it just repeats the single section total already shown above.
    if has_receivable and has_payable:
        ws.cell(row=row, column=1, value="GENEL NET BAKİYE").font = Font(size=12, bold=True)
        net_value = ledger.balance.net_balance
        label = "Bakiye Yok" if abs(net_value) < 0.005 else ("Müşteri Borcu" if net_value > 0 else "Bizim Borcumuz")
        net_cell = ws.cell(row=row, column=7, value=f"{abs(net_value):,.2f} TRY ({label})" if abs(net_value) >= 0.005 else "0,00 TRY (Bakiye Yok)")
        net_cell.font = Font(size=12, bold=True)
        net_cell.alignment = RIGHT

    widths = [14, 12, 18, 40, 16, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    safe_name = "".join(c for c in ledger.company.name if c.isalnum() or c in (" ", "-", "_")).strip()
    filename = f"{safe_name}_cari_hesap_{dt.date.today().isoformat()}.xlsx"
    file_path = os.path.join(settings.export_dir, filename)
    wb.save(file_path)
    return file_path
